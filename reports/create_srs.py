from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'srs_assets'
ASSETS.mkdir(exist_ok=True)
SOURCE = Path(r'C:\Users\alok0\Downloads\BTech_3rd_Year_SRS_Template.docx')
OUT = ROOT / 'MarketLab_SRS.docx'
doc = Document(SOURCE)
paragraphs = list(doc.paragraphs)
tables = list(doc.tables)
TITLE = 'MarketLab: Historical Stock Market Simulator with Machine Learning'
BLUE = '#173d4d'
TEAL = '#087f78'

def font(size=24):
    return ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', size)

def canvas(w=1500, h=800):
    im = Image.new('RGB', (w, h), 'white')
    return im, ImageDraw.Draw(im)

def label(draw, xy, text, size=24, color=BLUE):
    draw.multiline_text(xy, text, font=font(size), fill=color, align='center', anchor='mm', spacing=7)

def box(draw, bounds, text, ellipse=False, size=24):
    if ellipse:
        draw.ellipse(bounds, fill='#eff8f5', outline=TEAL, width=3)
    else:
        draw.rounded_rectangle(bounds, radius=16, fill='#eff8f5', outline=TEAL, width=3)
    label(draw, ((bounds[0]+bounds[2])/2, (bounds[1]+bounds[3])/2), text, size)

def arrow(draw, start, end, text='', offset=(0,-28)):
    import math
    draw.line([start,end], fill=BLUE, width=3)
    angle=math.atan2(end[1]-start[1],end[0]-start[0])
    pts=[end,(end[0]-17*math.cos(angle-.5),end[1]-17*math.sin(angle-.5)),(end[0]-17*math.cos(angle+.5),end[1]-17*math.sin(angle+.5))]
    draw.polygon(pts,fill=BLUE)
    if text: label(draw,((start[0]+end[0])/2+offset[0],(start[1]+end[1])/2+offset[1]),text,20)

def save(im, name):
    path=ASSETS/name
    im.save(path)
    return path

# Context diagram and architecture are generated as actual figures, not placeholders.
im,d=canvas(1500,650)
box(d,(35,215,330,365),'Student / learner')
box(d,(560,185,970,395),'MarketLab\nHistorical trading and\ndecision review')
box(d,(1150,55,1465,205),'Historical dataset\nImport stage')
box(d,(1150,410,1465,560),'SQLite storage\nPrices and user records')
arrow(d,(330,260),(560,260),'Orders and predictions')
arrow(d,(560,335),(330,335),'Charts and results',offset=(0,30))
arrow(d,(1150,160),(970,225),'OHLCV data',offset=(0,-25))
arrow(d,(970,360),(1150,450),'Persist and retrieve',offset=(20,-30))
context=save(im,'context.png')

im,d=canvas(1500,650)
for b,t in [((40,210,370,390),'Presentation layer\nNext.js, React, TypeScript\nTailwind and charts'),((570,210,930,390),'Application layer\nFastAPI and Python\nTrading rules and NumPy ML'),((1130,210,1460,390),'Data layer\nSQLAlchemy and SQLite\nPrices, trades and models')]: box(d,b,t,size=23)
arrow(d,(370,255),(570,255),'HTTP / JSON')
arrow(d,(570,345),(370,345),'Responses',offset=(0,30))
arrow(d,(930,255),(1130,255),'Queries')
arrow(d,(1130,345),(930,345),'Records',offset=(0,30))
label(d,(750,90),'MarketLab application architecture',32)
architecture=save(im,'architecture.png')

im,d=canvas(1500,850)
d.rounded_rectangle((420,35,1455,810),radius=20,outline=BLUE,width=3)
label(d,(940,75),'MarketLab system boundary',28)
box(d,(30,345,285,485),'Registered learner',size=23)
cases=[('Create a simulation',(530,115,940,205)),('Research prices and model',(1000,235,1400,325)),('Place and cancel orders',(530,355,940,445)),('Advance historical time',(1000,475,1400,565)),('Record a prediction',(530,595,940,685)),('Review results and lessons',(1000,695,1400,785))]
for text,b in cases:
    box(d,b,text,ellipse=True,size=21)
    d.line([(285,415),(b[0],(b[1]+b[3])/2)],fill='#83979d',width=2)
usecase=save(im,'usecase.png')

im,d=canvas(1500,900)
box(d,(30,350,270,510),'Learner')
nodes=[((400,55,750,195),'1. Authenticate\nand manage sessions'),((1000,55,1450,195),'2. Import and retrieve\nhistorical data'),((400,345,750,485),'3. Process orders\nand advance replay'),((1000,345,1450,485),'4. Train and evaluate\nmarket model'),((400,670,750,810),'5. Record and resolve\njournal predictions'),((1000,670,1450,810),'6. Present portfolio\nand results')]
for b,t in nodes: box(d,b,t,size=23)
arrow(d,(270,385),(400,125),'Login',offset=(-20,-25))
arrow(d,(270,430),(400,415),'Orders')
arrow(d,(270,470),(400,740),'Calls',offset=(-10,15))
arrow(d,(1000,145),(750,365),'Observed prices',offset=(30,-20))
arrow(d,(1225,195),(1225,345),'Past features',offset=(95,0))
arrow(d,(1000,450),(750,700),'Saved signal',offset=(20,-20))
arrow(d,(750,750),(1000,750),'Outcomes')
arrow(d,(750,460),(1000,700),'Fills and snapshots',offset=(30,-10))
label(d,(750,865),'Persistent stores: users, sessions, prices, orders, fills, models, journal and snapshots',21)
dfd=save(im,'dfd.png')

im,d=canvas(1500,920)
entities=[((30,30,420,150),'User\nPK id'),((555,30,945,150),'Simulation\nPK id | FK user_id'),((1080,30,1470,150),'Instrument\nPK id | unique symbol'),((30,315,420,455),'Order\nPK id | FK simulation_id\nFK instrument_id'),((555,315,945,455),'Holding / Snapshot\nFK simulation_id\nHolding: FK instrument_id'),((1080,315,1470,455),'DailyPrice\nPK id | FK instrument_id\nUnique instrument and date'),((30,680,420,820),'Fill\nPK id | unique FK order_id'),((555,680,945,820),'JournalEntry\nFK simulation_id, instrument_id\nOptional unique FK order_id'),((1080,680,1470,820),'ModelRun\nFK simulation_id, instrument_id\nVersion and training cutoff')]
for b,t in entities: box(d,b,t,size=22)
arrow(d,(420,90),(555,90),'1 to many')
arrow(d,(640,150),(330,315),'1 to many',offset=(-30,-10))
arrow(d,(750,150),(750,315),'1 to many',offset=(95,0))
arrow(d,(1275,150),(1275,315),'1 to many',offset=(95,0))
arrow(d,(225,455),(225,680),'1 to 0..1',offset=(95,0))
d.line([(945,100),(1010,100),(1010,595),(850,595)],fill=BLUE,width=3)
arrow(d,(850,595),(850,680),'',offset=(0,0))
d.line([(1010,595),(1250,595)],fill=BLUE,width=3)
arrow(d,(1250,595),(1250,680),'',offset=(0,0))
label(d,(1115,565),'Simulation: 1 to many',21)
label(d,(750,875),'Additional stores: CashLedger, Watchlist, ImportRun and Challenge.\nArrows show selected relationships; each model and journal record also belongs to a simulation.',21)
erd=save(im,'erd.png')

im,d=canvas(1500,960)
xs=[180,570,950,1320]
for x,t in zip(xs,['Learner','Browser UI','FastAPI','SQLite']):
    box(d,(x-130,35,x+130,125),t,size=25)
    d.line([(x,125),(x,890)],fill='#98abb1',width=2)
for y,a,b,t in [(200,0,1,'Place buy order'),(270,1,2,'POST order with token'),(340,2,3,'Validate ownership, cash and save Pending'),(410,3,2,'Saved order'),(480,2,1,'Pending confirmation'),(560,0,1,'Click Next day'),(630,1,2,'POST advance'),(700,2,3,'Fill at next open; update cash and shares'),(770,3,2,'Commit fill and equity snapshot'),(840,2,1,'Updated date and portfolio')]:
    arrow(d,(xs[a],y),(xs[b],y),t,offset=(0,-23))
label(d,(750,925),'Invalid quantity, insufficient cash or an ended session returns an error before a fill is created.',21)
sequence=save(im,'sequence.png')

im,d=canvas(1500,850)
for left,title in [(30,'Login'),(530,'Replay workspace'),(1030,'Results')]:
    d.rounded_rectangle((left,30,left+440,805),radius=16,outline=BLUE,width=3)
    label(d,(left+220,75),title,28)
    if left==30:
        box(d,(left+35,210,left+405,290),'Email')
        box(d,(left+35,340,left+405,420),'Password')
        box(d,(left+35,490,left+405,570),'Log in / Register')
    elif left==530:
        box(d,(left+25,140,left+415,220),'Current date | Next day',size=22)
        box(d,(left+25,260,left+415,445),'Historical chart')
        box(d,(left+25,480,left+415,620),'Stock | Buy / Sell\nQuantity | Submit',size=22)
        box(d,(left+25,660,left+415,760),'Pending / filled orders',size=22)
    else:
        box(d,(left+25,140,left+415,270),'Equity | Benchmark | Risk',size=22)
        box(d,(left+25,310,left+415,490),'Equity curve')
        box(d,(left+25,530,left+415,650),'Your call | Model | Actual',size=22)
        box(d,(left+25,685,left+415,760),'Trade review',size=22)
wireframe=save(im,'wireframes.png')

replacements={
1:'DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\nABES Engineering College, Ghaziabad',
3:f'Project Title: {TITLE}\nProject Code: P_195\nFormat reference: IEEE Std 830-1998',
4:'',
6:'Document Revision and Review History',
11:'This document describes the requirements and current scope of MarketLab. It explains how a learner can study historical Indian stock prices, place virtual orders, record a prediction and review the outcome. It is intended for the project guide, academic evaluators and the student developer. It also provides a reference for future changes to the application.',
13:'Beginners often judge a trading decision only after they know the result. MarketLab provides a practical way to record the decision first and then reveal what happened. The application replays historical prices from January 2019 to December 2021 and allows users to practise with virtual money.',
14:'In scope: account registration and login, historical data import, simulation setup, stock research, virtual trading, automatic model training, decision journaling, portfolio tracking, learning lessons, a 20-day challenge and results comparison.',
15:'Outside the current scope: real money trading, live brokerage orders, intraday execution, derivatives, short selling, leverage, an administrator dashboard and downloadable PDF or spreadsheet reports.',
16:'Expected benefits: learners can practise without risking money, compare their predictions with recorded outcomes and understand the difference between a correct direction forecast and a profitable trade. These are learning goals, rather than measured claims of financial improvement.',
19:'IEEE Std 830-1998, IEEE Recommended Practice for Software Requirements Specifications. Used as the structural reference supplied with the college template.',
20:'Ian Sommerville, Software Engineering, 10th edition, Pearson. General reference for software requirements and system design.',
21:'Project source and README: https://github.com/GuTS805/MarketLab. Implementation references: backend/app/main.py, backend/app/ml.py, backend/app/models.py and backend/import_market.py. Dataset: vishnun0027/indian-market-historical-ohlcv, revision fe783082c7ec68847ecd0d0b4e1ca06364507e5f.',
23:'Section 2 describes the application, users and operating environment. Section 3 lists the functional requirements. Section 4 explains the interfaces. Section 5 records quality requirements and verification methods. Section 6 presents the main design models. Section 7 provides space for academic review and signatures.',
26:'MarketLab is a browser-based application with a Next.js presentation layer, a FastAPI application layer and a SQLite data layer accessed through SQLAlchemy. Historical prices are imported before a demonstration. The browser requests only information available through the current simulation date. The same backend handles trading, model evaluation and journal outcomes.',
29:'The application supports visitors and registered learners. The developer maintains the local installation and imports the dataset through a command-line script. There is no separate administrator account or administrative web console in the current application.',
31:'The project is intended for local development and academic demonstration. The current setup uses Windows, Python, Node.js and a modern browser.',
32:'Client: a modern browser with JavaScript enabled. The interface adapts to desktop and mobile widths. A keyboard or touch input is sufficient for ordinary use.',
33:'Runtime: Python 3.10 or later for the backend, subject to dependency compatibility, and Node.js 20.9 or later for Next.js 16. The application has been tested locally with Python 3.13. GPU hardware is not required for the NumPy model.',
34:'Storage: SQLite is the default database. A practical development machine should have enough disk space for dependencies, imported historical rows and saved simulation records. A Windows laptop with 8 GB RAM is a reasonable recommended setup, not a measured minimum.',
36:'Budget and data: the implementation uses open source libraries and a public historical dataset. The imported date range and selected stocks define which simulations can be run.',
37:'Security: passwords are stored as bcrypt hashes. Signed JWT tokens identify users, and backend ownership checks restrict access to each user\'s simulations. Logout removes the browser token. Server-side revocation and token expiry are not implemented in the current version.',
38:'Design: page components, API access, trading rules, database models and machine learning code are kept in separate files. Prices use fixed-point database values and cash is rounded to two decimal places. Daily OHLC data limits the precision of simulated order execution.',
40:'Assumptions: learners understand that the money is virtual and that historical outcomes do not guarantee future results. The local database remains available while the application is running. Simulation time is controlled by the saved session date, rather than the device clock.',
41:'Dependencies: Internet access is needed to install packages and download the dataset, unless local copies are supplied. After import, the core local simulation does not depend on a live market feed or continuous access to the dataset host.',
43:'The following requirements describe the working application. Each requirement has a unique identifier so that implementation and verification can be traced to a specific feature.',
44:'3.1 Module 1: Accounts and Simulation Access',
45:'3.2 Module 2: Historical Data, Trading and Machine Learning',
46:'3.3 Module 3: Learning, Review and Visualization',
49:'The interface provides a login screen, simulation setup, an overview, stock research, replay, portfolio, Model Insights, a decision journal, results and learning pages. Forms display validation errors and action feedback. Completed simulations show final balances and a new-simulation option. A closed portfolio displays sale history instead of suggesting that no trading has taken place.',
52:'MarketLab uses ordinary computer input and display devices.',
53:'Standard peripherals: keyboard, mouse or touch input, and a browser display. Charts support stock selection and interactive navigation.',
54:'No sensors, embedded controllers or dedicated trading hardware are required. The model runs on the CPU.',
56:'Browser to server: Next.js and React call FastAPI endpoints through a shared fetch helper. Pydantic validates inputs. SQLAlchemy connects the backend to SQLite.',
57:'Data and charts: Pandas and PyArrow read historical Parquet data, NumPy trains the model, Lightweight Charts renders stock candles and volume, and Recharts renders portfolio and evaluation graphs. Lucide React provides interface icons.',
59:'Local traffic: the frontend normally runs on localhost:3000 and the API on localhost:8000. The launcher can select free ports when these are occupied. Local development uses HTTP. A public deployment would require an HTTPS configuration.',
60:'Requests and responses generally use UTF-8 JSON. Login uses form-encoded credentials. Authenticated requests carry a Bearer token in the Authorization header. CORS allows the configured localhost development origins.',
63:'The diagrams below describe the current design. The context and data flow diagrams show the application boundary and main processes. The relationship diagram summarizes the main database links. The sequence diagram follows a buy order from submission to its next-day fill.',
64:'7. Academic Review and Sign-Off',
65:'Prepared by Alok Srivastava for review under the guidance of Mr. Faizan. Signatures and review dates are to be entered by the relevant academic reviewers after evaluation. No approval is implied by this document.'
}
for i,text in replacements.items(): paragraphs[i].text=text

def fill(table, rows):
    prototype=deepcopy(table.rows[-1]._tr)
    while len(table.rows)>len(rows): table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows)<len(rows): table._tbl.append(deepcopy(prototype))
    for row,values in zip(table.rows,rows):
        for cell,value in zip(row.cells,values): cell.text=value

fill(tables[0],[[
    'SUBMITTED BY\nAlok Srivastava\nRoll Number: 2400320100149\n\nDegree: Bachelor of Technology (B.Tech)\nYear: Third Year\nSemester: Fifth Semester\nAcademic Year: 2026-27',
    'UNDER THE GUIDANCE OF\nMr. Faizan\nDepartment of Computer Science and Engineering\n\nProject: P_195\nDocument: Software Requirements Specification\nDate of Submission: 1 October 2026'
]])
fill(tables[1],[['Version','Date','Description','Prepared By','Review Status'],['1.0','1 October 2026','Project-specific SRS prepared for academic submission','Alok Srivastava','Pending academic review']])
fill(tables[2],[['Document note: This report follows the section structure of the supplied B.Tech SRS template. It describes the implemented project and identifies current limits. Performance and security statements are not presented as independently certified results.']])
fill(tables[3],[['Category','Term','Meaning'],['Documentation','SRS','Software Requirements Specification'],['Interface','API / REST','Application endpoints used to exchange requests and responses'],['Security','JWT','Signed token used to identify a logged-in user'],['Market data','OHLCV','Open, high, low, close and trading volume'],['Trading','P&L','Profit and loss, either realised or unrealised'],['Accounting','Equity','Cash plus the latest value of open holdings'],['ML','Logistic regression','Classifier that estimates the probability of an Up outcome'],['Evaluation','Holdout / baseline','Later test examples / a simple comparison prediction'],['Risk','Drawdown','Decline in portfolio value from an earlier peak']])
fill(tables[5],[['User / Role','Experience','Responsibilities and Access'],['Registered learner','Basic to intermediate','Create simulations, research stocks, place virtual orders, record predictions and review personal results.'],['Visitor','Basic','View the login page and register or log in. Protected simulation data requires authentication.'],['Developer / maintainer','Technical','Install dependencies, run the importer, maintain the database and verify code. This is an operating responsibility, not an in-app admin role.']])
fill(tables[6],[['Req ID','Feature Description','Input / Validation Rule','Priority'],['FR-1.1','Register a learner account','Valid email; password length from 8 to 72 characters; duplicate email is rejected.','High'],['FR-1.2','Authenticate an existing account','Email and password are checked against the stored bcrypt hash.','High'],['FR-1.3','Protect personal simulation data','A valid signed token and matching session ownership are required.','High'],['FR-1.4','Create a simulation','Choose a starting date, positive cash balance and optional trading costs. Non-trading dates move to the next available date.','High'],['FR-1.5','End and reopen a saved review','Ending a session blocks new trades and date advances; saved results remain readable.','High']])
fill(tables[7],[['Req ID','Feature Description','Input / Validation Rule','Priority'],['FR-2.1','Import reproducible historical data','Import 12 stocks and NIFTY 50 for 2019 to 2021 from a pinned revision or local dataset folder.','High'],['FR-2.2','Validate imported records','Reject missing columns, null values, duplicate dates, invalid OHLC bounds and inadequate coverage.','High'],['FR-2.3','Restrict revealed information','Market queries, model inputs and journal outcomes shall not reveal data after the simulation date.','High'],['FR-2.4','Place virtual market and limit orders','Validate stock, Buy/Sell side, positive integer quantity, active session and positive limit price when required.','High'],['FR-2.5','Reserve cash and shares','Pending buys reserve estimated value plus charges; pending sells reserve shares. Reject insufficient availability.','High'],['FR-2.6','Cancel a pending order','Only a pending order owned by the learner can be cancelled.','Medium'],['FR-2.7','Advance historical time','Reveal the next available trading day or a selected group of trading days. Stop when no later data exists.','High'],['FR-2.8','Execute eligible orders','Market orders fill at a later open. Limits use that day\'s open and high/low bounds. Same-day execution is excluded.','High'],['FR-2.9','Maintain average-cost holdings','Update quantity and average cost on purchases; calculate realised P&L on sales. Partial sales preserve average cost.','High'],['FR-2.10','Account for charges and corporate actions','Apply enabled simulated fees; credit dividends and apply imported split ratios during date advances.','High'],['FR-2.11','Save equity snapshots','Persist cash and close-marked equity for every revealed trading day.','High'],['FR-2.12','Prepare historical features','Calculate six return, volatility, range and volume features from observed raw OHLCV only. Exclude split-crossing windows.','High'],['FR-2.13','Train models automatically','When a signal is requested, train on available history if no run exists for that session, stock, version and date. Require 160 labelled examples.','High'],['FR-2.14','Evaluate the direction classifier','Use older 80% for evaluation training and later 20% for holdout testing, with three earlier walk-forward checks and a majority baseline.','High'],['FR-2.15','Save and reuse a model run','Persist preprocessing statistics, weights, cutoff, sample counts, test metrics and held-out predictions.','High'],['FR-2.16','Return a next-day direction score','Apply sigmoid; classify Up at probability 0.5 or above. Explain missing history and model quality.','High'],['FR-2.17','Personalize decision feedback','Use resolved one-day journal calls for the same stock and simulation to report correct Up/Down calls. Do not claim these records retrain market weights.','Medium']])
fill(tables[8],[['Req ID','Feature Description','Input / Validation Rule','Priority'],['FR-3.1','Show stock and portfolio charts','Display historical candles and volume, equity snapshots, cash, allocation and risk measures.','High'],['FR-3.2','Record a decision before its outcome','Save stock, reason, Up/Down prediction and a 1, 5 or 20-day horizon; validate the reason and optional linked order.','High'],['FR-3.3','Freeze comparable model calls','For a one-day journal entry, save available model and baseline predictions. Later model updates shall not replace the original record.','High'],['FR-3.4','Resolve journal outcomes','Reveal the actual direction only after the required number of later stock closes is available.','High'],['FR-3.5','Review closed positions','When all holdings are sold, display closed-position status, cash balance and completed sale history.','High'],['FR-3.6','Compare decisions in Results','Show personal calls, saved model calls and actual outcomes; compare model and baseline on the same one-day records.','High'],['FR-3.7','Compare portfolio performance','Calculate return, NIFTY 50 comparison, maximum drawdown, available volatility and post-trade price movement.','High'],['FR-3.8','Provide learning and a challenge','Offer short lessons and a 20-trading-day challenge. Freeze its completed debrief at the challenge endpoint.','Medium'],['FR-3.9','Maintain a personal watchlist','Add or remove imported stocks for the authenticated user.','Medium'],['FR-3.10','Show historical context and ranking','Display curated events only through the revealed date; group leaderboard entries by identical simulation starting date.','Medium']])
fill(tables[10],[['Attribute','Requirement / Acceptance Criterion','Verification Method'],['5.1 Performance','Ordinary pages shall provide loading feedback. Model results for an unchanged stock, date and version shall reuse the saved run. Cold training time must be measured separately; no concurrency target has been verified.','Repeat requests and inspect saved run IDs; measure local response times before assigning a performance target.'],['5.2 Security','Passwords shall be hashed; invalid authentication shall be rejected; one learner shall not access another learner\'s sessions, orders or journal. Token expiry, revocation and rate limiting remain future hardening work.','Authentication and ownership tests, database inspection and code review.'],['5.3 Reliability','Saved records shall survive a normal restart. Cash and holdings shall stay consistent with fills. Unique fill and snapshot keys shall help prevent duplicate execution.','Trading engine tests, restart checks and inspection of fills, ledger entries and snapshots.'],['5.4 Usability and portability','Forms shall explain failures and show order state. Completed-session screens shall provide suitable navigation. Layouts shall adapt to common desktop and mobile widths.','Browser walkthroughs and viewport checks; broader cross-browser testing remains to be completed.'],['5.5 Numerical correctness','Cash shall be rounded to paise; holdings shall use average-cost accounting; returns and risk measures shall use saved observations.','Tests for fills, partial sales, charges and portfolio totals.'],['5.6 ML reproducibility','Identical observed inputs shall produce repeatable training. Feature and test data shall follow chronological cutoffs. Insufficient history shall not produce a trained model.','Feature, deterministic training, chronological evaluation and API cutoff tests.']])
fill(tables[11],[['Model','Purpose','Included Figure'],['Context / architecture','Show the learner, historical data, application layers and storage.','Figures 1 and 2'],['Use cases','Show the registered learner\'s principal actions.','Figure 4'],['Data flow','Show authentication, price retrieval, trading, modeling and review.','Figure 5'],['Entity relationships','Show primary data entities and selected key relationships.','Figure 6'],['Sequence','Follow an order from submission to its next-day fill.','Figure 7']])
fill(tables[12],[['Project Guide','Internal Examiner','Head of Department'],['Mr. Faizan\nSignature: __________________\nDate: __________________','Name: __________________\nSignature: __________________\nDate: __________________','Name: __________________\nSignature: __________________\nDate: __________________']])

def figure_cell(table,path,caption):
    cell=table.cell(0,0)
    cell.text=''
    p=cell.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(6.0))
    p=cell.add_paragraph(caption)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic=True

figure_cell(tables[4],context,'Figure 1. System context for a learner using MarketLab.')
figure_cell(tables[9],wireframe,'Figure 3. Simplified wireframes of login, replay and results screens.')

def before(anchor,text,style=None):
    p=anchor.insert_paragraph_before(text,style)
    return p

anchor=paragraphs[28]
p=before(anchor,'')
p.add_run().add_picture(str(architecture),width=Inches(6.4))
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
before(anchor,'Figure 2. Presentation, application and data layers.').alignment=WD_ALIGN_PARAGRAPH.CENTER

anchor=paragraphs[47]
before(anchor,'3.4 Model Training and Decision Feedback','Heading 2')
before(anchor,'MarketLab uses a binary logistic regression classifier implemented with NumPy. Its task is to estimate whether the next available trading day closes above the current close. The model does not estimate an exact price or trade return. An unchanged close belongs to the Down class.')
before(anchor,'The six inputs are the 1-day, 5-day and 10-day returns, 10-day return volatility, the daily high-low range relative to the close, and volume relative to its previous 20-day average. Historical examples are labelled Up when the next observed close is higher. The current day has no known next-day label and is used only for prediction.')
before(anchor,'Training requires at least 160 valid labelled examples, normally 181 daily rows before exclusions. Features are standardized using statistics from the relevant training subset. Batch gradient descent runs for 360 epochs with learning rate 0.08 and L2 regularization 0.025. A sigmoid converts the weighted feature score into an Up probability.')
before(anchor,'The earlier 80% of labelled examples train the evaluation model; the later 20% form a chronological holdout. Three earlier expanding-window checks provide additional comparisons. The app reports accuracy, precision, recall, Brier score, confusion counts and loss history. It labels a model above baseline only when its final holdout and earlier checks meet the application\'s thresholds. After evaluation, a final model is fitted on all currently observed labels.')
before(anchor,'Training is automatic when Explore or Portfolio requests signals, when a specific model is viewed, or when a one-day journal entry needs a model snapshot. It runs during the request; there is no separate training scheduler. Saved runs are reused for the same session, stock, date and model version.')
before(anchor,'User-specific feedback uses past resolved one-day journal predictions for the same stock and simulation. It reports directional accuracy and the learner\'s tendency to choose Up or Down. At least five resolved calls are needed for pattern feedback. The current classifier is trained on OHLCV; the personal history does not retrain its weights or change the market probability.')
before(anchor,'3.5 Demonstration and Acceptance Flow','Heading 2')
before(anchor,'Create the Buying the Bottom scenario on 23 March 2020 with Rs. 1,00,000 and leave the demo-portfolio option unchecked. Research RELIANCE and record a one-day prediction. Place a market buy for 10 shares and confirm that it is Pending. Advance one trading day and check the fill, cash deduction, holdings and resolved journal outcome. Later, sell the shares and advance again. After ending the session, verify the final cash, closed-position message, sale history and Results comparison. Profit is an observed outcome, not an acceptance condition.')

anchor=paragraphs[64]
for heading,path,caption in [('6.1 Learner Use Cases',usecase,'Figure 4. Main interactions available to a registered learner.'),('6.2 Data Flow',dfd,'Figure 5. Level 1 view of the principal application processes.'),('6.3 Entity Relationships',erd,'Figure 6. Simplified relationship map of the main database entities.'),('6.4 Buy Order Sequence',sequence,'Figure 7. A market order is queued first and executed on a later trading day.')]:
    h=before(anchor,heading,'Heading 2')
    h.paragraph_format.page_break_before=True
    p=before(anchor,'')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(6.4))
    before(anchor,caption).alignment=WD_ALIGN_PARAGRAPH.CENTER
before(anchor,'Database structure: prices are separated from instruments, orders from their fills, and simulations from per-day snapshots. A model run has a unique session, stock, version and cutoff combination. A fill has a unique order link. A holding has a unique session and stock combination. Ledger entries retain cash movements for review. Figure 6 is a summary and does not draw every foreign key.')
before(anchor,'Sequence exceptions: an invalid token, an order outside the learner\'s session, an invalid quantity, insufficient available cash or an ended simulation produces a validation error. A market order may be rejected at execution if the actual opening cost exceeds available cash. A limit order remains pending until its condition is met or the learner cancels it.')
before(anchor,'6.5 Verification Status','Heading 2')
before(anchor,'At preparation, all 11 backend tests passed and the frontend production build completed with TypeScript checks. Existing tests cover next-day fills, average cost, partial sales, limits, reservations, ownership, model feature cutoffs, chronological evaluation, repeatable training, journal timing and challenge behaviour. These checks do not establish load capacity, certified security or future market performance.')

# Consistent academic formatting while retaining the supplied template's sections and tables.
normal=doc.styles['Normal']
normal.font.name='Calibri'
normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.12
for style_name,size in [('Heading 1',16),('Heading 2',12)]:
    s=doc.styles[style_name]
    s.font.name='Calibri'
    s.font.size=Pt(size)
    s.font.color.rgb=RGBColor.from_string('173D4D')
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(12)
    s.paragraph_format.space_after=Pt(7)
for p in doc.paragraphs:
    if re.match(r'^[1-7]\. ',p.text):
        p.style='Heading 1'
        p.paragraph_format.page_break_before=True
    elif re.match(r'^\d\.\d ',p.text): p.style='Heading 2'
    elif p.text=='Document Revision and Review History':
        p.style='Heading 1'
        p.paragraph_format.page_break_before=True
for i in [1,2,3]:
    p=paragraphs[i]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name='Calibri'
        r.font.size=Pt(14 if i!=2 else 23)
        r.bold=i!=3
        r.font.color.rgb=RGBColor.from_string('173D4D')

for table in doc.tables:
    table.autofit=True
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(5)
                p.paragraph_format.line_spacing=1.05
                for run in p.runs:
                    run.font.name='Calibri'
                    run.font.size=Pt(9.5 if len(table.columns)>2 else 10.5)
                    if ri==0 and len(table.rows)>1:
                        run.bold=True
                        run.font.color.rgb=RGBColor.from_string('173D4D')
        if len(table.rows)>1:
            props=row._tr.get_or_add_trPr()
            if ri==0:
                repeat=OxmlElement('w:tblHeader')
                props.append(repeat)
            no_split=OxmlElement('w:cantSplit')
            props.append(no_split)

for section in doc.sections:
    section.header.paragraphs[0].text='MarketLab | Software Requirements Specification | P_195'
    section.header.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in section.header.paragraphs[0].runs: r.font.size=Pt(8)
    footer=section.footer.paragraphs[0]
    footer.text='Alok Srivastava | 2400320100149 | Page '
    field=OxmlElement('w:fldSimple')
    field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in footer.runs: r.font.size=Pt(8)

doc.core_properties.title=TITLE
doc.core_properties.author='Alok Srivastava'
doc.core_properties.subject='B.Tech Third Year, Fifth Semester Software Requirements Specification'
doc.core_properties.keywords='MarketLab, P_195, SRS, historical trading, logistic regression'
doc.core_properties.comments=''
doc.save(OUT)

# Check for forgotten template text and punctuation requested by the student.
texts=[p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells]
joined='\n'.join(texts)
for bad in ['\u2014','\u2013','<Insert','<Project Title>','Placeholder','XXXXXX','YOLOv8','Cloudinary','RBAC']:
    if bad in joined: raise RuntimeError(f'Unresolved text: {bad}')
print(f'Created: {OUT}')
print(f'Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; embedded figures: {len(doc.inline_shapes)}')
